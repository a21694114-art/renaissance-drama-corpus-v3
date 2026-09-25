#!/usr/bin/env python3
"""08b_apply_review.py — carry Grace's decisions from topic_sheet.xlsx into the files the pipeline reads.

    python 08b_apply_review.py --runs <runs dir> --seed 42 [--xlsx <topic_sheet.xlsx>] [--drafts <csv>]

Everything downstream (09_aggregate, 11_figures, 12_map, 13_site, 14_genre_figures) reads
topics_B_s<seed>/topic_sheet.csv, not the workbook.  This script reads the 'topics' sheet of the
workbook, takes the review columns — Label, Notes, use_in_genre_analysis, pattern_basis, basis,
review_status, chunks_read, draft_label — and

  1. writes them into topic_sheet.csv (all other columns untouched);
  2. writes them into ab_spelling/drafts/topic_drafts_B_s<seed>.csv, so that re-running
     `run_b.sh sheet` later rebuilds the workbook WITH the decisions instead of the AI drafts.

review_status is derived (always relative to the AI drafts, so re-running is safe) unless Grace typed
a status of her own:
  user_confirmed — use_in_genre_analysis is included or contextual_only AND she touched the row
                   (changed the use away from the AI draft, or wrote a Label or Notes): the decision is hers;
  user_pending   — she set use_in_genre_analysis to pending (looked at it, undecided);
  user_labelled  — she wrote a Label but left the draft's use unchanged (name accepted, inclusion not decided);
  otherwise the AI status stays (draft_ai / checked_ai).
Label controls display only; inclusion is decided solely by use_in_genre_analysis.
Guards before anything is written: use values in {included, candidate, contextual_only, pending};
topic ids unique; the workbook's id set identical to topic_sheet.csv's; the workbook's size column equal
to the run's (a workbook from another run/seed is refused). A pre-review snapshot
(topic_sheet.pre_review.csv, drafts …pre_review.csv) is written once and never overwritten.
"""
import argparse, csv, shutil
from pathlib import Path

USE_VALUES = {'included', 'candidate', 'contextual_only', 'pending'}
REVIEW_COLS = ['draft_label', 'pattern_basis', 'use_in_genre_analysis', 'basis', 'review_status', 'chunks_read', 'Label', 'Notes']


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--runs', required=True); ap.add_argument('--seed', type=int, default=42); ap.add_argument('--variant', default='B')
    ap.add_argument('--xlsx', default=''); ap.add_argument('--drafts', default='')
    a = ap.parse_args()
    d = Path(a.runs) / f'topics_{a.variant}_s{a.seed}'
    xlsx = Path(a.xlsx) if a.xlsx else d / 'topic_sheet.xlsx'
    dpath = Path(a.drafts) if a.drafts else Path(__file__).resolve().parent.parent / 'drafts' / f'topic_drafts_{a.variant}_s{a.seed}.csv'

    import openpyxl
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    ws = wb['topics'] if 'topics' in wb.sheetnames else wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True); header = [str(h).strip() if h is not None else '' for h in next(it)]
    col = {h: i for i, h in enumerate(header)}
    missing = [c for c in ['topic'] + REVIEW_COLS if c not in col]
    if missing: raise SystemExit(f'workbook lacks columns {missing}')
    edited = {}
    for row in it:
        if row is None or row[col['topic']] in (None, ''): continue
        t = int(row[col['topic']])
        edited[t] = {c: ('' if row[col[c]] is None else str(row[col[c]]).strip()) for c in REVIEW_COLS}

    old_drafts = {}
    if dpath.exists():
        old_drafts = {int(r['topic']): r for r in csv.DictReader(open(dpath, encoding='utf-8'))}

    bad = {t: e['use_in_genre_analysis'] for t, e in edited.items() if e['use_in_genre_analysis'] and e['use_in_genre_analysis'] not in USE_VALUES}
    if bad: raise SystemExit(f'use_in_genre_analysis must be one of {sorted(USE_VALUES)}; found {bad}')

    # ---- guards: ids unique, same set as the run's sheet, same sizes (= same run)
    sp = d / 'topic_sheet.csv'
    rows = list(csv.DictReader(open(sp, encoding='utf-8'))); fields = list(rows[0].keys())
    ids_csv = [int(r['topic']) for r in rows]
    ids_wb = []
    ws2 = wb['topics'] if 'topics' in wb.sheetnames else wb[wb.sheetnames[0]]
    for row in ws2.iter_rows(min_row=2, values_only=True):
        if row is not None and row[col['topic']] not in (None, ''): ids_wb.append(int(row[col['topic']]))
    dup = sorted({t for t in ids_wb if ids_wb.count(t) > 1})
    if dup: raise SystemExit(f'refused: topic ids appear more than once in the workbook: {dup}')
    missing = sorted(set(ids_csv) - set(ids_wb)); extra = sorted(set(ids_wb) - set(ids_csv))
    if missing or extra: raise SystemExit(f'refused: workbook ids differ from {sp.name} — missing {missing}, unknown {extra}')
    if 'size' in col:
        size_csv = {int(r['topic']): str(r.get('size', '')).strip() for r in rows}
        wrong = []
        for row in ws2.iter_rows(min_row=2, values_only=True):
            if row is None or row[col['topic']] in (None, ''): continue
            t = int(row[col['topic']]); v = row[col['size']]
            if v is not None and str(int(v) if isinstance(v, float) else v).strip() != size_csv.get(t, ''): wrong.append(t)
        if wrong: raise SystemExit(f'refused: cluster sizes differ from the current run for topics {wrong[:10]} — is this workbook from another seed/run?')

    # ---- review_status, derived against the AI drafts (the pre-review snapshot once it exists, so that a
    #      second `review` run does not demote decisions made in the first: "touched" always means
    #      "differs from what the AI proposed", not "differs from last time")
    snap_d = dpath.with_name(dpath.stem + '.pre_review.csv')
    ai_drafts = {int(r['topic']): r for r in csv.DictReader(open(snap_d, encoding='utf-8'))} if snap_d.exists() else old_drafts
    KNOWN = ('', 'draft_ai', 'checked_ai', 'user_confirmed', 'user_pending', 'user_labelled')
    counts = {}
    for t, e in edited.items():
        ai = ai_drafts.get(t, {}); ai_use = ai.get('use_in_genre_analysis', '') or ''
        use_now = e['use_in_genre_analysis']; touched = bool(e['Label'] or e['Notes'] or (use_now and use_now != ai_use))
        if e['review_status'] in KNOWN:            # anything else = a status Grace typed herself: kept as is
            if use_now in ('included', 'contextual_only') and touched: e['review_status'] = 'user_confirmed'
            elif use_now == 'pending' and touched: e['review_status'] = 'user_pending'
            elif e['Label'] and use_now == ai_use: e['review_status'] = 'user_labelled'
            else: e['review_status'] = ai.get('review_status', '') or e['review_status'] or 'draft_ai'
        counts[e['review_status'] or '(blank)'] = counts.get(e['review_status'] or '(blank)', 0) + 1

    # ---- snapshots (once) and backups (each run)
    snap = sp.with_name('topic_sheet.pre_review.csv')
    if not snap.exists(): shutil.copy(sp, snap)
    if dpath.exists() and not dpath.with_name(dpath.stem + '.pre_review.csv').exists(): shutil.copy(dpath, dpath.with_name(dpath.stem + '.pre_review.csv'))
    shutil.copy(sp, sp.with_suffix('.csv.bak'))

    # 1. topic_sheet.csv
    for r in rows:
        e = edited.get(int(r['topic']))
        if e:
            for c in REVIEW_COLS:
                if c in r: r[c] = e[c]
    with open(sp, 'w', newline='', encoding='utf-8') as f:
        w = csv.DictWriter(f, fieldnames=fields); w.writeheader(); w.writerows(rows)

    # 2. drafts file (now carrying the decisions)
    dpath.parent.mkdir(parents=True, exist_ok=True)
    if dpath.exists(): shutil.copy(dpath, dpath.with_suffix('.csv.bak'))
    with open(dpath, 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f); w.writerow(['topic'] + REVIEW_COLS)
        for t in sorted(edited): w.writerow([t] + [edited[t][c] for c in REVIEW_COLS])

    from collections import Counter
    uses = Counter(e['use_in_genre_analysis'] or 'unclassified' for e in edited.values())
    print(f'{len(edited)} topics read from {xlsx.name}; labels filled: {sum(1 for e in edited.values() if e["Label"])}')
    print('use_in_genre_analysis:', ', '.join(f'{k} {v}' for k, v in uses.most_common()))
    print('review_status:', ', '.join(f'{k} {v}' for k, v in sorted(counts.items())))
    print(f'→ {sp}\n→ {dpath}  (backups *.csv.bak; first-time snapshots *.pre_review.csv)')


if __name__ == '__main__':
    main()
