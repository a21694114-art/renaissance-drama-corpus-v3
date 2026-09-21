#!/usr/bin/env python3
"""08_topic_sheet.py (v2, 2026-09-21) — naming / classification workbook for one topic run.

  python 08_topic_sheet.py --chunks <chunks dir> --runs <runs dir> --seed 42 \
        [--model Qwen/Qwen3-Embedding-0.6B] [--drafts <topic_drafts_B_s42.csv>] [--out <xlsx>]

Topic ids are the FINAL BERTopic ids of doc_topics.csv (0 = largest); the raw HDBSCAN label of
each topic is carried in column hdbscan_label (sheet id_map) — 05_diagnostics.py tables that were
keyed by raw labels must be read through that map.

Keywords: the count matrix (one document per topic = its cleaned member chunks concatenated;
CountVectorizer min_df=1, max_df=1.0, token [a-zA-Z]{2,}, the 03_topics stopword list) is weighted
with the installed bertopic.vectorizers.ClassTfidfTransformer (default settings: idf =
log(1 + avg_words_per_topic / total_count_of_word)), i.e. the same class-TF-IDF BERTopic itself
uses.  Top 30 words per topic are the candidates for the KeyBERT-style list (cosine of the word
embedding to the topic centroid, same embedding model) and the MMR list (lambda 0.5).  Words
capitalised in >= 80 % of their corpus occurrences (probable names, but also God/Lord…) carry a *.

Per topic: size (chunks, words), n_works, n_works_ge3 (works contributing >= 3 chunks),
dominant work and share, top-3-work share, top works, top author string and share (raw author
field, informative only), genre_mix_of_topic (direction topic -> genre; NOT genre -> topic),
play_type mix, three representative chunks nearest the centroid, and the classification columns:
draft_label / pattern_basis / use_in_genre_analysis / basis / review_status / chunks_read — filled
from --drafts when given (an AI draft file; review_status stays draft_ai until Grace changes it),
plus empty Label / Notes for Grace.

Sheets: topics, coverage (chunks/words/works per use-category and per genre, all chunks accounted
for incl. outliers), seed_match (each topic's best-overlapping clusters in the other seeds, by
shared chunks — never by id), summary (run stats, seed agreement, keyword settings), outliers
(assignment rate per genre), id_map.  topic_sheet.csv holds the topics sheet.
"""
import argparse, csv, json, re
from collections import Counter, defaultdict
from pathlib import Path
import numpy as np

STOP_RE = re.compile(r'\b[a-z]+\b')
CAP_RE = re.compile(r"\b[A-Za-z][a-z']+\b")
DRAFT_COLS = ['draft_label', 'pattern_basis', 'use_in_genre_analysis', 'basis', 'review_status', 'chunks_read']


def load_stop():
    from sklearn.feature_extraction.text import ENGLISH_STOP_WORDS
    import importlib.util
    here = Path(__file__).resolve().parent
    spec = importlib.util.spec_from_file_location('t03', here / '03_topics.py')
    m = importlib.util.module_from_spec(spec); spec.loader.exec_module(m)
    return set(ENGLISH_STOP_WORDS) | m.STOPWORDS


def class_tfidf_top(docs_by_topic, stop, k=30):
    """Standard BERTopic class-TF-IDF (bertopic.vectorizers.ClassTfidfTransformer, defaults) on a
    uniform vectorizer.  Returns {topic: [(word, score), ...]} and a settings dict."""
    from sklearn.feature_extraction.text import CountVectorizer
    from bertopic.vectorizers import ClassTfidfTransformer
    import bertopic
    topics = sorted(docs_by_topic)
    corpus = [' '.join(docs_by_topic[t]) for t in topics]
    vec = CountVectorizer(ngram_range=(1, 1), min_df=1, max_df=1.0, token_pattern=r'(?u)\b[a-zA-Z]{2,}\b', stop_words=list(stop))
    X = vec.fit_transform(corpus)
    tr = ClassTfidfTransformer(bm25_weighting=False, reduce_frequent_words=False)
    C = tr.fit(X).transform(X).tocsr()
    words = vec.get_feature_names_out()
    out = {}
    for i, t in enumerate(topics):
        row = C.getrow(i)
        idx = row.indices[np.argsort(-row.data)[:k]]
        vals = np.sort(row.data)[::-1][:k]
        out[t] = [(words[j], float(v)) for j, v in zip(idx, vals) if v > 0]
    settings = {'bertopic_version': bertopic.__version__, 'transformer': 'bertopic.vectorizers.ClassTfidfTransformer(bm25_weighting=False, reduce_frequent_words=False)',
                'vectorizer': 'CountVectorizer(min_df=1, max_df=1.0, token_pattern=[a-zA-Z]{2,}, stop_words=03_topics.STOPWORDS+sklearn english)',
                'documents': 'one per FINAL topic id: lowercase alphabetic tokens of member chunks minus stopwords', 'vocabulary_size': int(X.shape[1])}
    return out, settings


def mmr(cand_vecs, centroid, k=10, lam=0.5):
    rel = cand_vecs @ centroid
    chosen = [int(np.argmax(rel))]
    while len(chosen) < min(k, len(rel)):
        sim_to_chosen = (cand_vecs @ cand_vecs[chosen].T).max(axis=1)
        score = lam * rel - (1 - lam) * sim_to_chosen
        score[chosen] = -np.inf
        chosen.append(int(np.argmax(score)))
    return chosen


def agreement(ta, tb):
    from sklearn.metrics import adjusted_rand_score
    common = sorted(set(ta) & set(tb))
    la = [ta[c] for c in common]; lb = [tb[c] for c in common]
    both = [(x, y) for x, y in zip(la, lb) if x != -1 and y != -1]
    A = {c for c in common if ta[c] != -1}; B = {c for c in common if tb[c] != -1}
    return (round(adjusted_rand_score(la, lb), 3), round(adjusted_rand_score([x for x, _ in both], [y for _, y in both]), 3) if both else None,
            round(len(A & B) / max(1, len(A | B)), 3))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--chunks', required=True); ap.add_argument('--runs', required=True)
    ap.add_argument('--seed', type=int, default=42); ap.add_argument('--variant', default='B')
    ap.add_argument('--model', default='Qwen/Qwen3-Embedding-0.6B', help="embedding model for word scoring; 'none' = c-TF-IDF only")
    ap.add_argument('--drafts', default='', help='CSV with topic + draft columns (AI draft); default ab_spelling/drafts/topic_drafts_<V>_s<seed>.csv if present')
    ap.add_argument('--n-candidates', type=int, default=30); ap.add_argument('--n-rep', type=int, default=3)
    ap.add_argument('--out', default='')
    a = ap.parse_args()
    runs = Path(a.runs); ch = Path(a.chunks); d = runs / f'topics_{a.variant}_s{a.seed}'
    out = Path(a.out) if a.out else d / 'topic_sheet.xlsx'
    stop = load_stop()

    meta = {r['chunk_id']: r for r in csv.DictReader(open(ch / 'chunk_meta.csv', encoding='utf-8'))}
    wlen = {r['chunk_id']: int(r['len_B_words']) for r in csv.DictReader(open(ch / 'chunk_map.csv', encoding='utf-8'))}
    texts = {r['chunk_id']: r['text'] for r in csv.DictReader(open(ch / f'chunks_{a.variant}.csv', encoding='utf-8'))}
    labels = {r['chunk_id']: int(r['topic']) for r in csv.DictReader(open(d / 'doc_topics.csv', encoding='utf-8'))}
    raw = {r['chunk_id']: int(r['hdbscan_label']) for r in csv.DictReader(open(d / 'hdbscan_labels.csv', encoding='utf-8'))} if (d / 'hdbscan_labels.csv').exists() else {}
    ids = [r['chunk_id'] for r in csv.DictReader(open(runs / f'embedding_ids_{a.variant}.csv', encoding='utf-8'))]
    emb = np.load(runs / f'embeddings_{a.variant}.npy').astype(np.float32)
    emb /= np.linalg.norm(emb, axis=1, keepdims=True) + 1e-9
    row_of = {c: i for i, c in enumerate(ids)}
    assert set(labels) == set(ids), 'doc_topics.csv and embedding ids differ'

    # id map final topic -> raw hdbscan label (must be one-to-one)
    id_map = defaultdict(Counter)
    for c, t in labels.items():
        if t != -1: id_map[t][raw.get(c, '')] += 1
    id_map = {t: (cnt.most_common(1)[0][0] if len(cnt) == 1 else f'AMBIGUOUS {dict(cnt)}') for t, cnt in id_map.items()}

    # names by capitalisation ratio
    cap = Counter(); tot = Counter()
    for t in texts.values():
        for m in CAP_RE.finditer(t):
            w = m.group(0); lw = w.lower(); tot[lw] += 1
            if w[0].isupper(): cap[lw] += 1
    def is_name(w): return tot[w] >= 5 and cap[w] / tot[w] >= 0.8

    members = defaultdict(list)
    for c, t in labels.items():
        if t != -1: members[t].append(c)
    docs_by_topic = {t: [' '.join(w for w in STOP_RE.findall(texts[c].lower()) if w not in stop) for c in m] for t, m in members.items()}
    tw, kw_settings = class_tfidf_top(docs_by_topic, stop, a.n_candidates)
    cands = {t: [w for w, _ in tw[t][:a.n_candidates]] for t in members}
    with open(d / 'top_words_classtfidf.csv', 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f); w.writerow(['topic', 'rank', 'word', 'class_tfidf', 'probable_name'])
        for t in sorted(tw):
            for r, (wd, sc) in enumerate(tw[t], 1): w.writerow([t, r, wd, round(sc, 6), int(is_name(wd))])

    wvec = {}
    if a.model.lower() != 'none':
        from sentence_transformers import SentenceTransformer
        import torch
        dev = 'mps' if torch.backends.mps.is_available() else ('cuda' if torch.cuda.is_available() else 'cpu')
        model = SentenceTransformer(a.model, device=dev)
        vocab = sorted({w for ws in cands.values() for w in ws})
        V = model.encode(vocab, batch_size=64, normalize_embeddings=True, convert_to_numpy=True, show_progress_bar=True)
        wvec = dict(zip(vocab, V))

    # drafts (AI) — optional
    drafts = {}
    dpath = Path(a.drafts) if a.drafts else Path(__file__).resolve().parent.parent / 'drafts' / f'topic_drafts_{a.variant}_s{a.seed}.csv'
    if dpath.exists():
        for r in csv.DictReader(open(dpath, encoding='utf-8')):
            drafts[int(r['topic'])] = r
        print(f'drafts loaded from {dpath}: {len(drafts)} topics')

    rows = []
    for t, m in sorted(members.items(), key=lambda kv: -len(kv[1])):
        works = Counter(meta[c]['work_id'] for c in m)
        title_of = {}
        for c in m: title_of.setdefault(meta[c]['work_id'], meta[c]['title'])
        dom_work, dom_n = works.most_common(1)[0]
        top3 = sum(n for _, n in works.most_common(3))
        authors = Counter(meta[c]['author'] for c in m)
        genres = Counter(meta[c]['genre_deep'] for c in m); ptypes = Counter(meta[c]['play_type_deep'] for c in m)
        cent = emb[[row_of[c] for c in m]].mean(axis=0); cent /= np.linalg.norm(cent) + 1e-9
        cw = cands[t]; names = [w for w in cw if is_name(w)]
        mark = lambda w: ('*' + w) if w in names else w
        kb = mm = ''
        if wvec and cw:
            cv = np.stack([wvec[w] for w in cw]); rel = cv @ cent
            kb = ', '.join(mark(cw[i]) for i in np.argsort(-rel)[:10])
            mm = ', '.join(mark(cw[i]) for i in mmr(cv, cent, k=10, lam=0.5))
        sims = emb[[row_of[c] for c in m]] @ cent
        reps = [m[i] for i in np.argsort(-sims)[:a.n_rep]]
        rep_txt = '\n\n'.join(f'[{c} — {meta[c]["title"]} ({meta[c]["author"]}, {meta[c]["year"]}); {meta[c]["genre_deep"]}] {texts[c][:350]}…' for c in reps)
        dr = drafts.get(t, {})
        rows.append({'topic': t, 'hdbscan_label': id_map.get(t, ''), 'size': len(m), 'words': sum(wlen[c] for c in m),
                     'n_works': len(works), 'n_works_ge3': sum(1 for n in works.values() if n >= 3),
                     'dominant_work': title_of[dom_work], 'dominant_work_share': round(dom_n / len(m), 3), 'top3_work_share': round(top3 / len(m), 3),
                     'top_works': '; '.join(f'{title_of[w][:40]} ({n})' for w, n in works.most_common(5)),
                     'top_author_raw': authors.most_common(1)[0][0][:40], 'top_author_share': round(authors.most_common(1)[0][1] / len(m), 3),
                     'genre_mix_of_topic': '; '.join(f'{g} {n / len(m):.0%}' for g, n in genres.most_common(3)),
                     'play_type_mix': '; '.join(f'{g[:30]} {n / len(m):.0%}' for g, n in ptypes.most_common(2)),
                     'ctfidf_top10': ', '.join(mark(w) for w in cw[:10]), 'keybert_top10': kb, 'mmr_top10': mm, 'names_in_top30': ', '.join(names),
                     'representative_chunks': rep_txt,
                     **{k: dr.get(k, '') for k in DRAFT_COLS}, 'Label': dr.get('Label', ''), 'Notes': dr.get('Notes', '')})   # Label/Notes survive a re-run once 08b wrote them into the drafts file
        if dr and not rows[-1]['review_status']: rows[-1]['review_status'] = 'draft_ai'

    # coverage by use-category (all chunks accounted for)
    use_of = {r['topic']: (r['use_in_genre_analysis'] or 'unclassified') for r in rows}
    cat_of_chunk = {c: ('outlier_hdbscan' if t == -1 else use_of[t]) for c, t in labels.items()}
    cats = ['included', 'candidate', 'contextual_only', 'pending', 'unclassified', 'outlier_hdbscan']
    genres_all = Counter(meta[c]['genre_deep'] for c in labels)
    main_genres = [g for g, _ in genres_all.most_common(12)]
    cov_rows = []
    for cat in cats:
        cs = [c for c in labels if cat_of_chunk[c] == cat]
        if not cs and cat != 'outlier_hdbscan': continue
        r = {'use_category': cat, 'topics': sum(1 for t in use_of if use_of[t] == cat) if cat != 'outlier_hdbscan' else '',
             'chunks': len(cs), 'chunk_share': round(len(cs) / len(labels), 4), 'words': sum(wlen[c] for c in cs),
             'works': len({meta[c]['work_id'] for c in cs})}
        gc = Counter(meta[c]['genre_deep'] for c in cs)
        for g in main_genres: r[f'{g} chunks'] = gc[g]; r[f'{g} share of genre'] = round(gc[g] / genres_all[g], 3)
        cov_rows.append(r)
    tot_row = {'use_category': 'ALL', 'topics': len(rows), 'chunks': len(labels), 'chunk_share': 1.0, 'words': sum(wlen.values()), 'works': len({m['work_id'] for m in meta.values()})}
    for g in main_genres: tot_row[f'{g} chunks'] = genres_all[g]; tot_row[f'{g} share of genre'] = 1.0
    cov_rows.append(tot_row)

    # seed matching by shared chunks
    seeds = sorted(int(p.name.split('_s')[1]) for p in runs.glob(f'topics_{a.variant}_s*') if (p / 'doc_topics.csv').exists())
    lab = {s: {r['chunk_id']: int(r['topic']) for r in csv.DictReader(open(runs / f'topics_{a.variant}_s{s}' / 'doc_topics.csv', encoding='utf-8'))} for s in seeds}
    match_rows = []
    for t, m in sorted(members.items(), key=lambda kv: -len(kv[1])):
        ms = set(m); r = {'topic': t, 'size': len(m)}
        for s in seeds:
            if s == a.seed: continue
            ov = Counter(lab[s][c] for c in m)
            outl = ov.pop(-1, 0)
            best = ov.most_common(3)
            parts = []
            for bt, n in best:
                sz = sum(1 for x in lab[s].values() if x == bt)
                parts.append(f'T{bt}: {n} shared / size {sz} (J={n / len(ms | {c for c, x in lab[s].items() if x == bt}):.2f})')
            r[f'seed{s}_best'] = ' ; '.join(parts); r[f'seed{s}_outlier_share'] = round(outl / len(m), 3)
            r[f'seed{s}_status'] = ('stable' if best and best[0][1] / len(m) >= 0.7 and best[0][1] / sum(1 for x in lab[s].values() if x == best[0][0]) >= 0.7
                                    else 'split' if len(best) >= 2 and best[1][1] / len(m) >= 0.2 else 'merged/other' if best and best[0][1] / len(m) >= 0.5 else 'dissolved')
        match_rows.append(r)

    summary = []
    for s in seeds:
        rj = json.load(open(runs / f'topics_{a.variant}_s{s}' / 'run.json'))
        mem = Counter(t for t in lab[s].values() if t != -1); wk = defaultdict(Counter)
        for c, t in lab[s].items():
            if t != -1: wk[t][meta[c]['work_id']] += 1
        sw = [t for t in mem if wk[t].most_common(1)[0][1] / mem[t] >= 0.9]
        summary.append({'seed': s, 'topics': rj['n_topics'], 'outlier_share': rj['outlier_share'], 'topics_dominant_work_ge90': len(sw), 'chunks_in_those': sum(mem[t] for t in sw)})
    pairs = []
    for i in range(len(seeds)):
        for j in range(i + 1, len(seeds)):
            ari_all, ari_both, jac = agreement(lab[seeds[i]], lab[seeds[j]])
            pairs.append({'seeds': f'{seeds[i]}-{seeds[j]}', 'ARI_all_chunks': ari_all, 'ARI_assigned_in_both (conditional)': ari_both, 'Jaccard_assigned_sets': jac})
    unassigned = Counter(meta[c]['genre_deep'] for c, t in labels.items() if t == -1)
    out_rows = [{'genre_deep': g, 'chunks': n, 'unassigned': unassigned[g], 'assignment_rate': round(1 - unassigned[g] / n, 3)} for g, n in genres_all.most_common()]

    # write
    with open(out.with_suffix('.csv'), 'w', newline='', encoding='utf-8') as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0])); w.writeheader(); w.writerows(rows)
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    wb = Workbook(); ws = wb.active; ws.title = 'topics'
    cols = list(rows[0]); ws.append(cols)
    for r in rows: ws.append([r[c] for c in cols])
    widths = {'topic': 7, 'hdbscan_label': 8, 'size': 7, 'words': 8, 'n_works': 7, 'n_works_ge3': 8, 'dominant_work': 26, 'dominant_work_share': 9, 'top3_work_share': 9,
              'top_works': 40, 'top_author_raw': 20, 'top_author_share': 9, 'genre_mix_of_topic': 28, 'play_type_mix': 28, 'ctfidf_top10': 38, 'keybert_top10': 38,
              'mmr_top10': 38, 'names_in_top30': 28, 'representative_chunks': 80, 'draft_label': 34, 'pattern_basis': 20, 'use_in_genre_analysis': 14,
              'basis': 40, 'review_status': 12, 'chunks_read': 16, 'Label': 24, 'Notes': 30}
    for i, c in enumerate(cols, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = widths.get(c, 12); ws.cell(row=1, column=i).font = Font(bold=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row: cell.alignment = Alignment(wrap_text=True, vertical='top')
    ws.freeze_panes = 'B2'
    def sheet(name, rws):
        s = wb.create_sheet(name)
        if rws:
            s.append(list(rws[0]))
            for r in rws: s.append([r.get(k, '') for k in rws[0]])
            for cell in s[1]: cell.font = Font(bold=True)
        return s
    sheet('coverage', cov_rows); sheet('seed_match', match_rows)
    s2 = wb.create_sheet('summary')
    s2.append(['run', str(runs.name), 'variant', a.variant, 'seed', a.seed, 'word-scoring model', a.model])
    s2.append(['drafts file', str(dpath) if dpath.exists() else '(none)', 'review_status of drafts', 'draft_ai = AI proposal, not confirmed by Grace'])
    s2.append([]); s2.append(['keyword settings'] + [f'{k}: {v}' for k, v in kw_settings.items()])
    s2.append([]); s2.append(list(summary[0]))
    for r in summary: s2.append(list(r.values()))
    if pairs:
        s2.append([]); s2.append(list(pairs[0]))
        for r in pairs: s2.append(list(r.values()))
        s2.append(['note', 'ARI_assigned_in_both is conditional on both seeds assigning the chunk; Jaccard shows how many chunks flip between assigned and outlier.'])
    sheet('outliers', out_rows)
    sheet('id_map', [{'topic': t, 'hdbscan_label': id_map[t]} for t in sorted(id_map)])
    wb.save(out)
    (d / 'topic_sheet_coverage.csv').write_text('\n'.join([','.join(cov_rows[0].keys())] + [','.join(str(v) for v in r.values()) for r in cov_rows]), encoding='utf-8')
    print(f'{len(rows)} topics → {out}')
    print('keyword settings:', kw_settings)
    print('seeds:', summary); print('agreement:', pairs)
    print('coverage:', [(r['use_category'], r['chunks'], r['chunk_share']) for r in cov_rows])


if __name__ == '__main__':
    main()
